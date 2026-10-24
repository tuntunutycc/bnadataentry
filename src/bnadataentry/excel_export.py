"""Export invoices from SQLite to client-shaped Excel workbooks."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from bnadataentry.clients import get_client
from bnadataentry.models import Invoice
from bnadataentry.repository import InvoiceRepository

PROJECT_ROOT = Path(__file__).resolve().parents[2]
EXPORTS_DIR = PROJECT_ROOT / "exports"


def _fmt_date(value) -> str:
    if value is None:
        return ""
    return f"{value.day}.{value.month}.{value.year}"


def client_workbook_bytes(
    repo: InvoiceRepository,
    client_id: str,
    *,
    year: int | None = None,
    year_from: int | None = None,
    year_to: int | None = None,
    month: int | None = None,
) -> bytes:
    """Client workbook in memory. Callers do not need a file on disk."""
    wb = _build_client_workbook(
        repo,
        client_id,
        year=year,
        year_from=year_from,
        year_to=year_to,
        month=month,
    )
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_client_workbook(
    repo: InvoiceRepository,
    client_id: str,
    output_path: Path | str | None = None,
    *,
    year: int | None = None,
    year_from: int | None = None,
    year_to: int | None = None,
    month: int | None = None,
) -> Path:
    payload = client_workbook_bytes(
        repo,
        client_id,
        year=year,
        year_from=year_from,
        year_to=year_to,
        month=month,
    )
    if output_path is None:
        client = get_client(client_id)
        EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
        suffix = ""
        if year_from is not None and year_to is not None:
            if year_from == year_to and month:
                suffix = f"_{year_from}-{month:02d}"
            elif year_from == year_to:
                suffix = f"_{year_from}"
            else:
                suffix = f"_{year_from}-{year_to}"
                if month:
                    suffix += f"-{month:02d}"
        elif year and month:
            suffix = f"_{year}-{month:02d}"
        elif year:
            suffix = f"_{year}"
        output_path = EXPORTS_DIR / f"{client.name} - BNN Only Sale{suffix}.xlsx"

    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(payload)
    return path


def _build_client_workbook(
    repo: InvoiceRepository,
    client_id: str,
    *,
    year: int | None = None,
    year_from: int | None = None,
    year_to: int | None = None,
    month: int | None = None,
) -> Workbook:
    client = get_client(client_id)
    if year is not None and year_from is None and year_to is None:
        year_from = year
        year_to = year
    invoices = repo.list_invoices(
        client_id=client_id,
        year_from=year_from,
        year_to=year_to,
        month=month,
    )

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    summary = wb.create_sheet(f"Summary - {client.name}")
    _write_summary_sheet(summary, client_id, invoices)

    by_branch: dict[str, list[Invoice]] = {}
    for inv in invoices:
        branch = inv.branch or (client.branches[0] if client.branches else "Detail")
        by_branch.setdefault(branch, []).append(inv)

    for branch in client.branches:
        by_branch.setdefault(branch, [])

    for branch, branch_invoices in by_branch.items():
        sheet_name = _sheet_name_for_branch(client_id, branch)
        ws = wb.create_sheet(sheet_name[:31])
        _write_detail_sheet(ws, client_id, branch_invoices)
    return wb


def _sheet_name_for_branch(client_id: str, branch: str) -> str:
    mapping = {
        ("capital", "CREL"): "CREL Distribution Center",
        ("capital", "DC MDY"): "DC MDY",
        ("capital", "Hyper"): "Hyper ",
        ("capital", "G&G"): "G&G",
        ("makro", "Sat San"): "Makro-Sat Sun",
        ("makro", "Kha Wae Chan"): "Makro - Kha Wei Chan",
        ("one_stop_mart", "One Stop Mart"): "One Stop Mart",
    }
    return mapping.get((client_id, branch), branch)


def _write_summary_sheet(
    ws: Worksheet, client_id: str, invoices: list[Invoice]
) -> None:
    if client_id == "capital":
        headers = [
            "Po Date",
            "Invoice Date",
            "Due Date",
            "Invoice No",
            "Product",
            "CREL",
            "DC MDY",
            "Hyper",
            "G & G",
            "Retro (3%)",
            "Subtotal",
            "DC Mdy6%, Other 4%, Hyper 0%",
            "Balance (To Receive)",
            "Remark",
        ]
        ws.append(headers)
        for inv in invoices:
            ba = inv.summary.branch_amounts or {}
            merchandise = sum((line.line_total or 0) for line in inv.lines)
            branch_sum = sum(float(ba.get(b, 0) or 0) for b in ("CREL", "DC MDY", "Hyper", "G&G"))
            if not branch_sum:
                branch_sum = merchandise
            retro = inv.summary.retro_amount or 0
            invoice_balance = float(inv.summary.balance_to_receive or 0)
            subtotal = branch_sum + float(retro)
            if not subtotal:
                subtotal = invoice_balance
            # Capital DC is applied at payment; export stored DC (0 if unset).
            dc = abs(float(inv.summary.delivery_charge or 0.0))
            excel_balance = round(invoice_balance - dc, 2)
            ws.append(
                [
                    _fmt_date(inv.po_date),
                    _fmt_date(inv.invoice_date),
                    _fmt_date(inv.due_date),
                    inv.invoice_no,
                    inv.product_label,
                    ba.get("CREL", 0),
                    ba.get("DC MDY", 0),
                    ba.get("Hyper", 0),
                    ba.get("G&G", 0),
                    retro,
                    subtotal,
                    -abs(dc),  # Excel DC column is negative
                    excel_balance,
                    inv.summary.remark,
                ]
            )
    elif client_id == "makro":
        headers = [
            "PO Date",
            "Invoice Date",
            "Due Date",
            "Invoice No",
            "Product Description",
            "Sat San",
            "Kha Wae Chan",
            "Tax 5%",
            "Grand Total",
            "Retro ( -5% )",
            "Amount",
            "Remark",
        ]
        ws.append(headers)
        for inv in invoices:
            ba = inv.summary.branch_amounts or {}
            merchandise = sum((line.line_total or 0) for line in inv.lines)
            tax = inv.summary.tax_amount or 0
            ws.append(
                [
                    _fmt_date(inv.po_date),
                    _fmt_date(inv.invoice_date),
                    _fmt_date(inv.due_date),
                    inv.invoice_no,
                    inv.product_label,
                    ba.get("Sat San", 0),
                    ba.get("Kha Wae Chan", 0),
                    tax,
                    merchandise + tax,
                    inv.summary.retro_amount or 0,
                    inv.summary.balance_to_receive or 0,
                    inv.summary.remark,
                ]
            )
    else:  # one_stop_mart
        headers = [
            "PO Date",
            "Invoice Date",
            "Due Date",
            "Invoice No",
            "Product",
            "Total Invoice",
            "Retro (5%,10%)",
            "Balance (To Receive)",
        ]
        ws.append(headers)
        for inv in invoices:
            merchandise = sum((line.line_total or 0) for line in inv.lines)
            ws.append(
                [
                    _fmt_date(inv.po_date),
                    _fmt_date(inv.invoice_date),
                    _fmt_date(inv.due_date),
                    inv.invoice_no,
                    inv.product_label,
                    merchandise,
                    inv.summary.retro_amount or 0,
                    inv.summary.balance_to_receive or 0,
                ]
            )


def _write_detail_sheet(
    ws: Worksheet, client_id: str, invoices: list[Invoice]
) -> None:
    if client_id == "capital":
        ws.append(
            [
                "Po No",
                "Po Date",
                "Invoice Date",
                "Due Date (Payment)",
                "Invoice No",
                "Product Description",
                "Pack Qty",
                "Unit",
                "List Price",
                "Total",
                "Discount 5%, 3%",
                "Grand Total",
            ]
        )
        ws.append(["", "", "", "", "", "", "Pk", "", "", "", "", ""])
    elif client_id == "makro":
        ws.append(
            [
                "Po Date",
                "Delivery Date",
                "Due Date (Payment)",
                "Invoice No",
                "Product Description",
                "Pack Qty",
                "Unit",
                "List Price",
                "Total Invoice",
                "Tax 5%",
                "Grand Total",
                "Remark",
            ]
        )
        ws.append(["", "", "", "", "", "ဖာ", "Pcs", "", "", "", "", ""])
    else:
        ws.append(
            [
                "Po Date",
                "Invoice Date",
                "Due Date (Payment)",
                "Invoice No",
                "Product Description",
                "Pack Qty",
                "Unit",
                "List Price",
                "Total",
                "Discount (5%,10%)",
                "Grand Total",
            ]
        )
        ws.append(["", "", "", "", "", "Pk (pcs)", "", "", "", "", ""])

    for inv in invoices:
        first = True
        for line in inv.lines:
            if client_id == "capital":
                row = [
                    inv.po_no if first else "",
                    _fmt_date(inv.po_date) if first else "",
                    _fmt_date(inv.invoice_date) if first else "",
                    _fmt_date(inv.due_date) if first else "",
                    inv.invoice_no if first else "",
                    line.product_description,
                    line.pack_qty,
                    "",
                    0 if line.is_foc else line.list_price,
                    "FOC" if line.is_foc else (line.line_total or 0),
                    line.discount_amount,
                    "FOC" if line.is_foc else (line.grand_total or 0),
                ]
            elif client_id == "makro":
                row = [
                    _fmt_date(inv.po_date) if first else "",
                    _fmt_date(inv.invoice_date) if first else "",
                    _fmt_date(inv.due_date) if first else "",
                    inv.invoice_no if first else "",
                    line.product_description,
                    line.pack_qty_fa if line.pack_qty_fa else "",
                    line.pack_qty,
                    0 if line.is_foc else line.list_price,
                    "FOC" if line.is_foc else (line.line_total or 0),
                    line.tax_amount,
                    "FOC" if line.is_foc else (line.grand_total or 0),
                    inv.summary.remark if first else "",
                ]
            else:
                row = [
                    _fmt_date(inv.po_date) if first else "",
                    _fmt_date(inv.invoice_date) if first else "",
                    _fmt_date(inv.due_date) if first else "",
                    inv.invoice_no if first else "",
                    line.product_description,
                    line.pack_qty,
                    "",
                    line.list_price,
                    line.line_total or 0,
                    line.discount_amount,
                    line.grand_total or 0,
                ]
            ws.append(row)
            first = False

        # Total row
        if client_id == "capital":
            ws.append(
                [
                    "Total",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    0,
                    sum((line.line_total or 0) for line in inv.lines),
                    sum(line.discount_amount for line in inv.lines),
                    sum((line.grand_total or 0) for line in inv.lines),
                ]
            )
        elif client_id == "makro":
            ws.append(
                [
                    "Total Amount",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    0,
                    sum((line.line_total or 0) for line in inv.lines),
                    sum(line.tax_amount for line in inv.lines),
                    sum((line.grand_total or 0) for line in inv.lines),
                    "",
                ]
            )
        else:
            ws.append(
                [
                    "Total",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    0,
                    sum((line.line_total or 0) for line in inv.lines),
                    sum(line.discount_amount for line in inv.lines),
                    sum((line.grand_total or 0) for line in inv.lines),
                ]
            )
