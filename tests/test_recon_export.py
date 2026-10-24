"""Excel export for the live reconciliation screen. No sample workbooks."""

from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from bnadataentry.db import init_db
from bnadataentry.expected_payment import money
from bnadataentry.recon_export import export_filename, reconciliation_workbook_bytes
from bnadataentry.reconciliation import (
    InvoiceCategory,
    Product,
    ReconciliationRepository,
    SupermarketInvoice,
    SupermarketPayment,
)

INVOICE_DATE = date(2026, 10, 1)


@pytest.fixture
def repo(tmp_path: Path) -> ReconciliationRepository:
    return ReconciliationRepository(init_db(tmp_path / "export.db"))


def _sheet(payload: bytes, name: str):
    book = load_workbook(BytesIO(payload))
    return book[name]


def test_export_filename_drops_punctuation() -> None:
    assert export_filename("Capital", "G&G") == "Capital-G-G.xlsx"
    assert export_filename("Makro", "Makro - Kha Wei Chan").endswith(".xlsx")
    assert " " not in export_filename("Makro", "Makro - Kha Wei Chan")


def test_empty_shop_workbook_has_headers_only(
    repo: ReconciliationRepository,
) -> None:
    payload = reconciliation_workbook_bytes(repo, "Capital", "DC MDY")
    book = load_workbook(BytesIO(payload))
    assert book.sheetnames == [
        "Reconciliation",
        "Lines",
        "Payments",
        "By category",
        "By product",
    ]
    sheet = book["Reconciliation"]
    assert sheet["A1"].value == "Invoice No."
    assert sheet["A1"].font.bold
    assert sheet.max_row == 1


def test_capital_workbook_contains_the_invoice_line_and_payment(
    repo: ReconciliationRepository,
) -> None:
    chips = repo.save_product(
        Product(
            category="Snacks",
            subcategory="Chips",
            name="Potato chips",
            pieces_per_pack=12,
            price_per_piece=10,
        )
    )
    saved = repo.create_invoice(
        SupermarketInvoice(
            invoice_number="T-CAP-1",
            supermarket_name="Capital",
            shop="Hyper",
            invoice_date=INVOICE_DATE,
            due_date=INVOICE_DATE + timedelta(days=30),
            total_invoice_amount=1,
            credit_days=30,
            categories=[
                InvoiceCategory(
                    category=chips.category,
                    subcategory=chips.subcategory,
                    product_name=chips.name,
                    unit="Pack",
                    entered_quantity=2,
                    pieces_per_pack=chips.pieces_per_pack,
                    price_per_piece=chips.price_per_piece,
                )
            ],
        )
    )
    assert saved.id is not None
    repo.create_payment(
        SupermarketPayment(
            invoice_id=saved.id,
            payment_date=date(2026, 10, 20),
            actual_received_amount=232.80,
            retro_percentage=0.03,
            delivery_charge_percentage=0.0,
        )
    )
    payload = reconciliation_workbook_bytes(repo, "Capital", "Hyper")
    recon = _sheet(payload, "Reconciliation")
    assert [cell.value for cell in recon[2]][:13] == [
        "T-CAP-1",
        "Capital",
        "Hyper",
        "2026-10-01",
        "2026-10-31",
        30,
        0,
        None,
        240,
        232.8,
        232.8,
        0,
        "Paid",
    ]
    line = _sheet(payload, "Lines")
    assert line["D2"].value == "Potato chips"
    assert line["E2"].value == "Pack"
    assert line["F2"].value == 2
    assert line["I2"].value == 24
    assert line["J2"].value == 240
    assert line["K2"].value == "No"
    payment = _sheet(payload, "Payments")
    assert payment["A2"].value == "T-CAP-1"
    assert payment["B2"].value == "2026-10-20"
    assert payment["C2"].value == 232.8
    assert payment["D2"].value == 3
    category = _sheet(payload, "By category")
    assert category["A2"].value == "Snacks"
    assert category["B2"].value == "Chips"
    assert category["C2"].value == 240
    product = _sheet(payload, "By product")
    assert product["A2"].value == "Potato chips"
    assert product["D2"].value == 24
    assert product["E2"].value == money(240)


def test_makro_export_keeps_tax_retro_and_received_date(
    repo: ReconciliationRepository,
) -> None:
    saved = repo.create_invoice(
        SupermarketInvoice(
            invoice_number="M-1",
            supermarket_name="Makro",
            shop="Makro-Sat Sun",
            invoice_date=INVOICE_DATE,
            due_date=INVOICE_DATE + timedelta(days=30),
            total_invoice_amount=100_000,
            tax_percentage=0.05,
            retro_percentage=0.05,
            discount_percentage=0.03,
            credit_days=30,
            categories=[InvoiceCategory(category="Snacks", quantity=1, amount=100_000)],
        )
    )
    assert saved.id is not None
    repo.create_payment(
        SupermarketPayment(
            invoice_id=saved.id,
            payment_date=date(2026, 10, 20),
            actual_received_amount=1,
            tax_percentage=0.05,
            retro_percentage=0.05,
            retro_cut_off_date=date(2026, 10, 15),
        )
    )
    recon = _sheet(
        reconciliation_workbook_bytes(repo, "Makro", "Makro-Sat Sun"),
        "Reconciliation",
    )
    headers = [cell.value for cell in recon[1]]
    row = {headers[index]: recon.cell(2, index + 1).value for index in range(len(headers))}
    assert row["Discount %"] == 3
    assert row["Remark"] == "Discount 3%: 3,000.00"
    assert row["Tax %"] == 5
    assert row["Retro %"] == 5
    assert row["Received date"] == "2026-10-20"
    assert row["Retro cut-off"] == "2026-10-15"
    net = 100_000 * 0.97
    expected = money(net * 1.05 * 0.95)
    assert row["Expected"] == expected
